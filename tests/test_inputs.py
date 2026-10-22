from datetime import datetime, timezone
import pytest
from fastapi.testclient import TestClient
from auto_plan.app import create_app
from auto_plan.input_models import InputSnapshot
from auto_plan.settings import Settings
from tests.support import CompanyAuthentication, MemoryAccessStore


@pytest.fixture
def inputs_client():
    from tests.input_support import MemoryInputStore
    access = MemoryAccessStore()
    access.users['EDITOR'] = {'user_id': 'EDITOR', 'roles': ['editor'], 'active': True}
    app = create_app(Settings(public_origin='https://testserver'), store=access,
                     input_store=MemoryInputStore(), authenticator=CompanyAuthentication(),
                     clock=lambda: datetime(2026, 10, 14, tzinfo=timezone.utc))
    with TestClient(app, base_url='https://testserver') as client:
        client.headers.update({'Origin': 'https://testserver', 'X-Auto-Plan-Request': '1'})
        assert client.post('/auto_plan/api/auth/login', json={'username': 'EDITOR', 'password': 'company-secret'}).status_code == 200
        yield client


def batch(client, changes, *, apply=True):
    version = client.get('/auto_plan/api/inputs').json()['input_version']
    payload = {'base_version': version, 'changes': changes}
    preview = client.post('/auto_plan/api/inputs/preview', json=payload)
    if apply and preview.status_code == 200 and preview.json()['valid']:
        return client.post('/auto_plan/api/inputs/apply', json=payload)
    return preview


def masters(client):
    result = batch(client, {'recipes': [{'code': 'R', 'name': '工艺', 'active': True}, {'code': 'S', 'name': '另一工艺', 'active': True}],
                            'lines': [{'code': 'L', 'name': '一线', 'active': True}],
                            'models': [{'code': 'M', 'name': '机种', 'recipe_code': 'R', 'active': True}]})
    assert result.status_code == 200


def test_editor_previews_without_writing_and_saves_fixed_recipe_master(inputs_client):
    client = inputs_client
    masters(client)
    version = client.get('/auto_plan/api/inputs').json()['input_version']
    preview = batch(client, {'models': [{'code': 'M', 'name': '改名', 'recipe_code': 'R', 'active': True}]}, apply=False)
    assert preview.json()['valid'] is True
    assert client.get('/auto_plan/api/inputs').json()['masters']['models'][0]['name'] == '机种'
    changed = batch(client, {'models': [{'code': 'M', 'name': '改名', 'recipe_code': 'S', 'active': True}]})
    assert changed.json()['valid'] is False
    assert changed.json()['errors'][0] == {'sheet': 'models', 'row': 2, 'field': 'recipe_code', 'reason': '机种的 Recipe 固定，不可更换'}
    assert client.get('/auto_plan/api/inputs').json()['input_version'] == version


def test_code_normalization_preserves_atomic_rejection_for_nontext_control_and_long_values(inputs_client):
    client=inputs_client
    masters(client)
    before=client.get('/auto_plan/api/inputs').json()
    for code in ([],{},12,None,'R\n','X'*65):
        checked=batch(client,{'recipes':[{'code':code,'name':'Invalid code','active':True}]},apply=False)
        assert checked.status_code==200
        assert checked.json()['valid'] is False
        assert any(row['field']=='code' for row in checked.json()['errors'])
    spaced=batch(client,{'recipes':[{'code':' '*10000+' R ','name':'工艺','active':True}]},apply=False)
    assert spaced.status_code==200 and spaced.json()['valid'] is True
    assert spaced.json()['changes']['recipes'][0]['code']=='R'
    assert client.get('/auto_plan/api/inputs').json()==before


def test_parameter_batch_is_atomic_strict_and_preserves_unsupplied_values(inputs_client):
    client = inputs_client
    masters(client)
    assert batch(client, {'capacities': [{'month': '2026-10', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 2400}],
                          'line_yields': [{'month': '2026-10', 'model_code': 'M', 'percent': '95.50'}],
                          'model_yields': [{'month': '2026-10', 'model_code': 'M', 'percent': '98'}],
                          'cycle_times': [{'model_code': 'M', 'days': '0'}],
                          'switch_durations': [{'from_recipe': 'R', 'to_recipe': 'S', 'hours': '0'}]}).status_code == 200
    before = client.get('/auto_plan/api/inputs').json()
    bad = {'capacities': [{'month': '2026-10', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 2600},
                          {'month': '2026-10', 'recipe_code': 'R', 'line_code': 'L', 'pieces': True}],
           'line_yields': [{'month': '2026-10', 'model_code': 'M', 'percent': '0'},
                           {'month': '2026-11', 'model_code': 'M', 'percent': '100.001'}],
           'cycle_times': [{'model_code': 'M', 'days': ''}]}
    payload = {'base_version': before['input_version'], 'changes': bad}
    preview = client.post('/auto_plan/api/inputs/preview', json=payload).json()
    assert preview['valid'] is False
    assert {(e['sheet'], e['row'], e['field']) for e in preview['errors']} >= {
        ('capacities', 3, 'pieces'), ('line_yields', 2, 'percent'), ('line_yields', 3, 'percent'), ('cycle_times', 2, 'days')}
    assert client.post('/auto_plan/api/inputs/apply', json=payload).status_code == 422
    assert client.get('/auto_plan/api/inputs').json() == before
    duplicate = batch(client, {'capacities': [{'month': '2026-10', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 0}] * 2})
    assert '重复' in duplicate.json()['errors'][0]['reason']
    assert batch(client, {'cycle_times': [{'model_code': 'M', 'days': '0.50'}]}).status_code == 200
    assert client.get('/auto_plan/api/inputs').json()['parameters']['capacities'][0]['pieces'] == 2400


def test_current_line_yield_input_requires_and_round_trips_model_identity(inputs_client):
    client = inputs_client
    masters(client)
    before = client.get('/auto_plan/api/inputs').json()
    missing_model = batch(client, {'line_yields': [{'month': '2026-10', 'percent': '95'}]}, apply=False)
    assert missing_model.status_code == 200
    assert missing_model.json()['valid'] is False
    assert missing_model.json()['errors'] == [{
        'sheet': 'line_yields',
        'row': 2,
        'field': 'model_code',
        'reason': '线体良率缺少机种 model_code，请使用包含机种的新格式或新模板',
    }]
    rejected = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': before['input_version'],
        'changes': {'line_yields': [{'month': '2026-10', 'percent': '95'}]},
    })
    assert rejected.status_code == 422
    assert client.get('/auto_plan/api/inputs?version=' + str(before['input_version'])).json() == before

    saved = batch(client, {'line_yields': [{'month': '2026-10', 'model_code': 'M', 'percent': '95'}]})
    assert saved.status_code == 200, saved.text
    assert saved.json()['parameters']['line_yields'] == [
        {'month': '2026-10', 'model_code': 'M', 'percent': '95'}]
    reloaded = client.get('/auto_plan/api/inputs').json()
    assert next(row for row in reloaded['parameters']['line_yields']
                if row['month'] == '2026-10' and row['model_code'] == 'M')['percent'] == '95'


def test_legacy_global_yields_expand_once_for_all_retained_models(inputs_client):
    client = inputs_client
    legacy = {
        'schema_version': 1,
        'input_version': 1,
        'masters': {
            'recipes': [{'code': 'R', 'name': '工艺', 'active': True}],
            'lines': [{'code': 'L', 'name': '一线', 'active': True}],
            'models': [
                {'code': 'A', 'name': '机种 A', 'recipe_code': 'R', 'active': True},
                {'code': 'B', 'name': '机种 B', 'recipe_code': 'R', 'active': True},
                {'code': 'C', 'name': '保留停用机种', 'recipe_code': 'R', 'active': False},
            ],
            'capabilities': [],
        },
        'parameters': {
            'line_yields': [{'month': '2026-10', 'percent': '95'},
                            {'month': '2026-10', 'model_code': 'B', 'percent': '80'},
                            {'month': '2026-11', 'percent': '90'}],
            'model_yields': [],
            'capacities': [],
            'cycle_times': [],
            'switch_durations': [],
        },
        'calendar': {'holidays': [], 'opening_exceptions': [], 'maintenance': [], 'experiments': []},
        'business': {'demands': [], 'confirmations': []},
    }
    store = client.app.state.input_store
    store.versions[1] = InputSnapshot.model_validate(legacy)
    store.version = 1

    current = client.get('/auto_plan/api/inputs')
    assert current.status_code == 200, current.text
    current = current.json()
    assert current['schema_version'] == 1
    assert current['input_version'] == 2
    values = {(row['month'], row['model_code']): row['percent']
              for row in current['parameters']['line_yields']}
    assert {key: value for key, value in values.items() if key[0] in {'2026-10', '2026-11'}} == {
        ('2026-10', 'A'): '95', ('2026-10', 'B'): '80', ('2026-10', 'C'): '95',
        ('2026-11', 'A'): '90', ('2026-11', 'B'): '90', ('2026-11', 'C'): '90',
    }
    assert all('model_code' in row for row in current['parameters']['line_yields'])
    assert client.get('/auto_plan/api/inputs?version=1').json() == legacy
    assert client.get('/auto_plan/api/inputs').json() == current

    edited = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': 2,
        'changes': {
            'models': [{'code': 'D', 'name': '新机种', 'recipe_code': 'R', 'active': True}],
            'line_yields': [{'month': '2026-10', 'model_code': 'A', 'percent': '96'}],
        },
    })
    assert edited.status_code == 200, edited.text
    reloaded = client.get('/auto_plan/api/inputs').json()
    reloaded_values = {(row['month'], row['model_code']): row['percent']
                       for row in reloaded['parameters']['line_yields']}
    assert reloaded_values['2026-10', 'A'] == '96'
    assert reloaded_values['2026-10', 'B'] == '80'
    assert reloaded_values['2026-10', 'C'] == '95'
    assert ('2026-10', 'D') not in reloaded_values
    assert client.get('/auto_plan/api/inputs?version=1').json() == legacy


def test_line_yield_errors_identify_duplicate_and_unknown_model_rows(inputs_client):
    client = inputs_client
    masters(client)
    before = client.get('/auto_plan/api/inputs').json()
    payload = {'base_version': before['input_version'], 'changes': {'line_yields': [
        {'month': '2026-10', 'model_code': 'M', 'percent': '95'},
        {'month': '2026-10', 'model_code': 'M', 'percent': '96'},
        {'month': '2026-10', 'model_code': 'UNKNOWN', 'percent': '90'},
    ]}}
    preview = client.post('/auto_plan/api/inputs/preview', json=payload)
    assert preview.status_code == 200
    assert preview.json()['valid'] is False
    assert {(row['row'], row['field']) for row in preview.json()['errors']} == {
        (3, 'month,model_code'), (4, 'model_code')}
    assert client.post('/auto_plan/api/inputs/apply', json=payload).status_code == 422
    assert client.get('/auto_plan/api/inputs?version=' + str(before['input_version'])).json() == before


def test_direct_apply_cannot_add_a_model_before_legacy_conversion(inputs_client):
    client = inputs_client
    legacy = InputSnapshot.model_validate({
        'schema_version': 1,
        'input_version': 1,
        'masters': {
            'recipes': [{'code': 'R', 'name': '工艺', 'active': True}],
            'lines': [{'code': 'L', 'name': '一线', 'active': True}],
            'models': [{'code': 'A', 'name': '机种 A', 'recipe_code': 'R', 'active': True}],
        },
        'parameters': {'line_yields': [{'month': '2026-10', 'percent': '95'}]},
    })
    store = client.app.state.input_store
    store.versions[1] = legacy
    store.version = 1

    stale = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': 1,
        'changes': {'models': [{'code': 'D', 'name': '新机种', 'recipe_code': 'R', 'active': True}]},
    })
    assert stale.status_code == 409
    converted = client.get('/auto_plan/api/inputs').json()
    assert [row['code'] for row in converted['masters']['models']] == ['A']
    assert [row for row in converted['parameters']['line_yields'] if row['month'] == '2026-10'] == [
        {'month': '2026-10', 'model_code': 'A', 'percent': '95'}]


def test_line_yield_template_is_current_but_pinned_legacy_export_keeps_its_shape(inputs_client):
    from io import BytesIO
    from openpyxl import load_workbook

    client = inputs_client
    template = load_workbook(BytesIO(client.get('/auto_plan/api/inputs/template.xlsx').content))
    assert [cell.value for cell in template['line_yields'][1]] == ['month', 'model_code', 'percent']

    legacy = InputSnapshot.model_validate({
        'schema_version': 1,
        'input_version': 1,
        'masters': {
            'recipes': [{'code': 'R', 'name': '工艺', 'active': True}],
            'models': [{'code': 'A', 'name': '机种 A', 'recipe_code': 'R', 'active': True}],
        },
        'parameters': {'line_yields': [{'month': '2026-10', 'percent': '95'}]},
    })
    store = client.app.state.input_store
    store.versions[1] = legacy
    store.version = 1
    exported = load_workbook(BytesIO(
        client.get('/auto_plan/api/inputs/export.xlsx?version=1').content))
    assert [cell.value for cell in exported['line_yields'][1]] == ['month', 'percent']
    assert [cell.value for cell in exported['line_yields'][2]] == ['2026-10', '95']


def test_availability_uses_model_month_yield_on_every_line_without_fallback(inputs_client):
    client = inputs_client
    masters(client)
    assert batch(client, {
        'lines': [{'code': 'L2', 'name': '二线', 'active': True}],
        'models': [{'code': 'N', 'name': '另一机种', 'recipe_code': 'R', 'active': True}],
        'line_yields': [{'month': '2026-10', 'model_code': 'M', 'percent': '95'}],
        'model_yields': [{'month': '2026-10', 'model_code': 'M', 'percent': '90'}],
    }).status_code == 200
    for line in ('L', 'L2'):
        result = client.get(
            f'/auto_plan/api/inputs/availability?month=2026-10&model_code=M&line_code={line}').json()
        assert result['line_yield'] == {'state': 'known', 'percent': '95'}
        assert result['model_yield'] == {'state': 'known', 'percent': '90'}
    other = client.get(
        '/auto_plan/api/inputs/availability?month=2026-10&model_code=N&line_code=L').json()
    assert other['line_yield'] == {'state': 'missing', 'percent': None}
    assert other['model_yield'] == {'state': 'missing', 'percent': None}


def test_missing_months_copy_previous_values_once_without_linking_or_overwriting(inputs_client):
    client = inputs_client
    masters(client)
    result = batch(client, {'capacities': [
        {'month': '2026-09', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 2400},
        {'month': '2026-12', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 3000}],
        'line_yields': [{'month': '2026-09', 'model_code': 'M', 'percent': '95'}]})
    assert result.status_code == 200
    data = client.get('/auto_plan/api/inputs').json()
    assert {r['month']: r['pieces'] for r in data['parameters']['capacities']} == {
        '2026-09': 2400, '2026-10': 2400, '2026-11': 2400, '2026-12': 3000, '2027-01': 3000, '2027-02': 3000}
    assert [r['month'] for r in data['parameters']['line_yields']] == ['2026-09', '2026-10', '2026-11', '2026-12', '2027-01', '2027-02']
    batch(client, {'capacities': [{'month': '2026-11', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 2600}]})
    data2 = client.get('/auto_plan/api/inputs').json()
    assert next(r['pieces'] for r in data2['parameters']['capacities'] if r['month'] == '2026-12') == 3000
    assert client.get(f"/auto_plan/api/inputs?version={data['input_version']}").json() == data


def test_calendar_allows_over_24_and_preserves_distinct_availability_states(inputs_client):
    client = inputs_client
    masters(client)
    payload = {'holidays': [{'date': '2026-10-17', 'holiday': True}],
               'opening_exceptions': [{'date': '2026-10-17', 'line_code': 'L', 'enabled': True}],
               'maintenance': [{'date': '2026-10-17', 'line_code': 'L', 'hours': 20}],
               'experiments': [{'date': '2026-10-17', 'line_code': 'L', 'hours': 10}]}
    assert batch(client, payload).json()['calendar']['experiments'][0]['hours'] == 10
    url = '/auto_plan/api/inputs/availability?month=2026-10&model_code=M&line_code=L&date=2026-10-17'
    missing = client.get(url).json()
    assert missing['capacity'] == {'state': 'missing', 'pieces': 0}
    assert missing['line_yield']['state'] == missing['model_yield']['state'] == missing['cycle_time']['state'] == 'missing'
    assert missing['switch'] == {'state': 'unknown_initial', 'hours': '0'}
    assert missing['available_hours'] == 0 and missing['output_holiday'] is True
    assert client.get(url + '&from_recipe=S').json()['switch'] == {'state': 'default_missing', 'hours': '12'}
    batch(client, {'capacities': [{'month': '2026-10', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 0}]})
    assert client.get(url).json()['capacity']['state'] == 'zero'
    batch(client, {'capabilities': [{'recipe_code': 'R', 'line_code': 'L', 'allowed': False}]})
    assert client.get(url).json()['capacity']['state'] == 'forbidden'
    old = client.get('/auto_plan/api/inputs').json()
    batch(client, {'models': [{'code': 'M', 'name': '历史机种', 'recipe_code': 'R', 'active': False}]})
    assert client.get('/auto_plan/api/inputs/options').json()['models'] == []
    assert client.get(f"/auto_plan/api/inputs?version={old['input_version']}").json() == old


def test_standard_excel_roundtrip_reports_source_rows_and_requires_atomic_apply(inputs_client):
    from io import BytesIO
    from openpyxl import load_workbook
    client = inputs_client
    masters(client)
    response = client.get('/auto_plan/api/inputs/template.xlsx')
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook['capacities'].cell(1, 1).value == 'month'
    workbook['capacities'].append(['2026-10', 'R', 'L', 2400])
    workbook['line_yields'].append(['2026-10', 'M', '95.50'])
    out = BytesIO()
    workbook.save(out)
    version = client.get('/auto_plan/api/inputs').json()['input_version']
    preview = client.post(f'/auto_plan/api/inputs/excel/preview?base_version={version}', content=out.getvalue(), headers={'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}).json()
    assert preview['valid'] is True
    assert client.get('/auto_plan/api/inputs').json()['parameters']['capacities'] == []
    assert client.post('/auto_plan/api/inputs/apply', json={'base_version': version, 'changes': preview['changes']}).status_code == 200
    export = client.get('/auto_plan/api/inputs/export.xlsx').content
    roundtrip = load_workbook(BytesIO(export))
    assert roundtrip['line_yields'].cell(2, 3).value == '95.50'
    # Wrong data on row 3 is reported at row 3, and an Excel formula is never evaluated.
    workbook['capacities'].append(['2026-11', 'R', 'L', '=1+1'])
    out = BytesIO(); workbook.save(out)
    version = client.get('/auto_plan/api/inputs').json()['input_version']
    preview = client.post(f'/auto_plan/api/inputs/excel/preview?base_version={version}', content=out.getvalue()).json()
    assert {'sheet': 'capacities', 'row': 3, 'field': 'pieces', 'reason': '不接受公式，请填写明确值'} in preview['errors']
    assert preview['valid'] is False


def test_multi_model_paste_export_and_reimport_preserve_unrelated_rows(inputs_client):
    from io import BytesIO
    from openpyxl import load_workbook

    client = inputs_client
    seeded = batch(client, {
        'recipes': [{'code': 'R', 'name': '工艺', 'active': True}],
        'lines': [{'code': 'L', 'name': '一线', 'active': True}],
        'models': [
            {'code': 'A', 'name': '机种 A', 'recipe_code': 'R', 'active': True},
            {'code': 'B', 'name': '机种 B', 'recipe_code': 'R', 'active': True},
        ],
        'capacities': [{'month': '2026-10', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 2400}],
        'model_yields': [
            {'month': '2026-10', 'model_code': 'A', 'percent': '90'},
            {'month': '2026-10', 'model_code': 'B', 'percent': '90'},
        ],
        'holidays': [{'date': '2026-10-17', 'holiday': True}],
    })
    assert seeded.status_code == 200, seeded.text
    before = client.get('/auto_plan/api/inputs').json()
    paste_preview = client.post('/auto_plan/api/inputs/paste/preview', json={
        'base_version': before['input_version'],
        'kind': 'line_yields',
        'text': 'month\tmodel_code\tpercent\n2026-10\tA\t95\n2026-10\tB\t80',
    })
    assert paste_preview.status_code == 200, paste_preview.text
    assert paste_preview.json()['valid'] is True
    assert client.get('/auto_plan/api/inputs').json() == before
    applied = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': paste_preview.json()['base_version'],
        'changes': paste_preview.json()['changes'],
    })
    assert applied.status_code == 200, applied.text
    saved = applied.json()

    exported_at_version = client.get(
        f"/auto_plan/api/inputs/export.xlsx?version={saved['input_version']}")
    assert exported_at_version.status_code == 200
    assert client.get(f"/auto_plan/api/inputs?version={saved['input_version']}").json() == saved
    workbook = load_workbook(BytesIO(exported_at_version.content))
    yields = workbook['line_yields']
    rows = {yields.cell(row, 2).value: row for row in range(2, yields.max_row + 1)}
    assert [yields.cell(1, column).value for column in range(1, 4)] == [
        'month', 'model_code', 'percent']
    assert rows.keys() == {'A', 'B'}
    yields.cell(rows['A'], 3).value = '96'
    edited = BytesIO()
    workbook.save(edited)

    excel_preview = client.post(
        f"/auto_plan/api/inputs/excel/preview?base_version={saved['input_version']}",
        content=edited.getvalue())
    assert excel_preview.status_code == 200, excel_preview.text
    assert excel_preview.json()['valid'] is True
    assert client.get(f"/auto_plan/api/inputs?version={saved['input_version']}").json() == saved
    reimported = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': excel_preview.json()['base_version'],
        'changes': excel_preview.json()['changes'],
    })
    assert reimported.status_code == 200, reimported.text
    current = reimported.json()

    assert current['parameters']['line_yields'] == [
        {'month': '2026-10', 'model_code': 'A', 'percent': '96'},
        {'month': '2026-10', 'model_code': 'B', 'percent': '80'},
    ]
    assert current['parameters']['model_yields'] == saved['parameters']['model_yields']
    assert current['parameters']['capacities'] == saved['parameters']['capacities']
    assert current['calendar']['holidays'] == saved['calendar']['holidays']
    before_export = client.get('/auto_plan/api/inputs').json()
    assert client.get('/auto_plan/api/inputs/export.xlsx').status_code == 200
    assert client.get('/auto_plan/api/inputs').json() == before_export


def test_multi_model_paste_errors_are_ordered_and_reject_the_whole_batch(inputs_client):
    client = inputs_client
    seeded = batch(client, {
        'recipes': [{'code': 'R', 'name': '工艺', 'active': True}],
        'models': [
            {'code': 'A', 'name': '机种 A', 'recipe_code': 'R', 'active': True},
            {'code': 'B', 'name': '机种 B', 'recipe_code': 'R', 'active': True},
        ],
    })
    assert seeded.status_code == 200, seeded.text
    before = client.get('/auto_plan/api/inputs').json()
    text = '\n'.join([
        'month\tmodel_code\tpercent',
        '2026-10\tA\t95',
        '2026-10\tUNKNOWN\t90',
        '2026-10\tA\t94',
        '2026-10\tB\t0',
    ])
    preview = client.post('/auto_plan/api/inputs/paste/preview', json={
        'base_version': before['input_version'], 'kind': 'line_yields', 'text': text})
    assert preview.status_code == 200, preview.text
    assert preview.json()['valid'] is False
    assert {(row['row'], row['field']) for row in preview.json()['errors']} == {
        (3, 'model_code'),
        (4, 'month,model_code'),
        (5, 'percent'),
    }

    rejected = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': before['input_version'],
        'changes': {'line_yields': [
            {'month': '2026-10', 'model_code': 'A', 'percent': '95'},
            {'month': '2026-10', 'model_code': 'UNKNOWN', 'percent': '90'},
            {'month': '2026-10', 'model_code': 'A', 'percent': '94'},
            {'month': '2026-10', 'model_code': 'B', 'percent': '0'},
        ]},
    })
    assert rejected.status_code == 422
    assert client.get('/auto_plan/api/inputs').json() == before

    corrected = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': before['input_version'],
        'changes': {'line_yields': [
            {'month': '2026-10', 'model_code': 'A', 'percent': '95'},
            {'month': '2026-10', 'model_code': 'B', 'percent': '80'},
        ]},
    })
    assert corrected.status_code == 200, corrected.text
    assert corrected.json()['parameters']['line_yields'] == [
        {'month': '2026-10', 'model_code': 'A', 'percent': '95'},
        {'month': '2026-10', 'model_code': 'B', 'percent': '80'},
    ]


def test_legacy_line_yield_excel_header_points_to_the_new_template(inputs_client):
    from io import BytesIO
    from openpyxl import Workbook

    client = inputs_client
    masters(client)
    before = client.get('/auto_plan/api/inputs').json()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'line_yields'
    sheet.append(['month', 'percent'])
    sheet.append(['2026-10', '95'])
    content = BytesIO()
    workbook.save(content)

    preview = client.post(
        f"/auto_plan/api/inputs/excel/preview?base_version={before['input_version']}",
        content=content.getvalue())

    assert preview.status_code == 200
    assert preview.json()['valid'] is False
    assert preview.json()['errors'] == [{
        'sheet': 'line_yields',
        'row': 1,
        'field': 'headers',
        'reason': '旧格式缺少机种列 model_code，请使用包含机种的新格式或新模板',
    }]
    assert client.get('/auto_plan/api/inputs').json() == before


def test_tabular_paste_uses_the_same_validation_and_keeps_missing_fields(inputs_client):
    client = inputs_client
    masters(client)
    before = client.get('/auto_plan/api/inputs').json()
    preview = client.post('/auto_plan/api/inputs/paste/preview', json={
        'base_version': before['input_version'], 'kind': 'maintenance',
        'text': 'date\tline_code\thours\n2026-10-17\tL\t20\n2026-10-18\tL\t'}).json()
    assert preview['valid'] is False
    assert preview['errors'][0]['row'] == 3 and preview['errors'][0]['field'] == 'hours'
    assert client.get('/auto_plan/api/inputs').json() == before


def test_legacy_line_yield_paste_header_points_to_the_new_format(inputs_client):
    client = inputs_client
    masters(client)
    before = client.get('/auto_plan/api/inputs').json()
    preview = client.post('/auto_plan/api/inputs/paste/preview', json={
        'base_version': before['input_version'], 'kind': 'line_yields',
        'text': 'month\tpercent\n2026-10\t95',})

    assert preview.status_code == 200
    assert preview.json()['valid'] is False
    assert preview.json()['errors'] == [{
        'sheet': 'line_yields',
        'row': 1,
        'field': 'headers',
        'reason': '旧格式缺少机种列 model_code，请使用包含机种的新格式或新模板',
    }]
    assert client.get('/auto_plan/api/inputs').json() == before


def test_stale_apply_cannot_overwrite_and_viewer_or_pure_admin_cannot_write(inputs_client):
    client = inputs_client
    masters(client)
    before = client.get('/auto_plan/api/inputs').json()
    batch(client, {'cycle_times': [{'model_code': 'M', 'days': '0.50'}]})
    stale = client.post('/auto_plan/api/inputs/apply', json={'base_version': before['input_version'], 'changes': {'cycle_times': [{'model_code': 'M', 'days': '2'}]}})
    assert stale.status_code == 409
    access = client.app.state.access_store
    for role in ['viewer', 'administrator']:
        access.users['EDITOR']['roles'] = [role]
        for path in ['/inputs/preview', '/inputs/apply']:
            assert client.post('/auto_plan/api' + path, json={'base_version': 0, 'changes': {}}).status_code == 403
        assert client.get('/auto_plan/api/inputs/template.xlsx').status_code == (200 if role == 'viewer' else 403)


def test_excel_percentage_cells_use_displayed_percent_units(inputs_client):
    from io import BytesIO
    from openpyxl import load_workbook
    client = inputs_client
    masters(client)
    book = load_workbook(BytesIO(client.get('/auto_plan/api/inputs/template.xlsx').content))
    sheet = book['line_yields']
    sheet.append(['2026-10', 'M', 0.95]); sheet['C2'].number_format = '0.00%'
    sheet.append(['2026-11', 'M', 1]); sheet['C3'].number_format = '0%'
    out = BytesIO(); book.save(out)
    version = client.get('/auto_plan/api/inputs').json()['input_version']
    result = client.post(f'/auto_plan/api/inputs/excel/preview?base_version={version}', content=out.getvalue()).json()
    assert result['valid'] is True
    assert result['changes']['line_yields'] == [
        {'month': '2026-10', 'model_code': 'M', 'percent': '95'},
        {'month': '2026-11', 'model_code': 'M', 'percent': '100'},
    ]


def test_availability_query_validates_month_date_and_the_production_month(inputs_client):
    client = inputs_client
    masters(client)
    root = '/auto_plan/api/inputs/availability?model_code=M&line_code=L'
    for query in ['&month=2026-00', '&month=2026-10&date=2026-10-33', '&month=2026-11&date=2026-10-17']:
        assert client.get(root + query).status_code == 422
