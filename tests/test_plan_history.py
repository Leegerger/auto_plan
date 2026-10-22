"""Selected history and restoration through the public business API."""
from datetime import datetime, timedelta, timezone
from tests.test_plan_transactions import shared_client, post, command, patch, row


def restore(identity=1, kind='revision'):
    return {'kind': 'restore_history', 'source': {'kind': kind, 'id': identity}}


def preview(client, request):
    return post(client, 'preview', {key: value for key, value in request.items() if key not in {'operation_id', 'pending'}})


def test_selected_history_keeps_original_inputs_results_and_missing_old_metrics(shared_client):
    client, plans, inputs = shared_client
    assert post(client, 'save', command(patch(row()))).status_code == 200
    original = client.get('/auto_plan/api/plans/revisions/1').json()
    # A legacy persisted metric is absent, never recomputed from today's context.
    plans.revisions[1]['evaluation'].pop('stability', None)
    changed = client.post('/auto_plan/api/inputs/apply', json={'base_version': 1, 'changes': {
        'capacities': [{'month': '2026-10', 'recipe_code': 'R', 'line_code': 'L', 'pieces': 2000}]
    }})
    assert changed.status_code == 200, changed.text
    selected = client.get('/auto_plan/api/plans/revisions/1').json()
    assert selected['snapshot'] == original['snapshot']
    assert 'stability' not in selected['evaluation']
    listing = client.get('/auto_plan/api/plans/history')
    assert listing.status_code == 200, listing.text
    assert [(r['revision'], r['actor'], r['recoverable']) for r in listing.json()['items']] == [(1, 'EDITOR', True)]
    assert client.get('/auto_plan/api/plans/current').json()['input_version'] == 2


def test_legacy_global_yields_remain_raw_in_history_after_current_masters_change(shared_client):
    import json
    from auto_plan.input_models import InputSnapshot
    client, _, inputs = shared_client
    legacy_document = inputs.read().model_dump(mode='json')
    legacy_document['parameters']['line_yields'] = [{'month': '2026-10', 'percent': '95'}]
    legacy_document['parameters']['model_yields'] = [
        {'month': '2026-10', 'model_code': 'M', 'percent': '98'}]
    legacy = InputSnapshot.model_validate(legacy_document)
    inputs.versions[1] = legacy
    inputs.version = 1
    saved = post(client, 'save', command(patch(row()), version=1))
    assert saved.status_code == 200, saved.text
    original_input = client.get('/auto_plan/api/inputs?version=1').json()
    original = client.get('/auto_plan/api/plans/revisions/1').json()
    assert original['snapshot']['inputs']['parameters']['line_yields'] == [{'month': '2026-10', 'percent': '95'}]

    converted = client.get('/auto_plan/api/inputs').json()
    changed = client.post('/auto_plan/api/inputs/apply', json={'base_version': converted['input_version'], 'changes': {
        'models': [{'code': 'Q', 'name': '新机种', 'active': True, 'recipe_code': 'R'}],
        'line_yields': [{'month': '2026-10', 'model_code': 'M', 'percent': '96'}]
    }})
    assert changed.status_code == 200, changed.text
    assert changed.json()['masters']['models'][-1]['code'] == 'Q'
    current_yields = {(row['month'], row['model_code']): row['percent']
                      for row in changed.json()['parameters']['line_yields']}
    assert current_yields['2026-10', 'M'] == '96'
    assert current_yields['2026-10', 'N'] == '95'
    assert ('2026-10', 'Q') not in current_yields

    assert client.get('/auto_plan/api/inputs?version=1').json() == original_input
    selected = client.get('/auto_plan/api/plans/revisions/1').json()
    assert selected == original
    rejected = client.post('/auto_plan/api/plans/evaluate', json={'snapshot': original['snapshot']})
    assert rejected.status_code == 422, rejected.text
    book = workbook(client.get('/auto_plan/api/plans/revisions/1/export.xlsx'))
    exported = json.loads(''.join(row['json_chunk'] for row in records(book, '_record_json')))
    assert exported == original


def test_legacy_history_keeps_original_output_while_restore_uses_current_model_yields(shared_client):
    import json
    from auto_plan.input_models import InputSnapshot
    client, _, inputs = shared_client
    legacy_document = inputs.read().model_dump(mode='json')
    legacy_document['parameters']['line_yields'] = [{'month': '2026-10', 'percent': '95'}]
    legacy_document['parameters']['model_yields'] = [
        {'month': '2026-10', 'model_code': 'M', 'percent': '90'}]
    inputs.versions[1] = InputSnapshot.model_validate(legacy_document)
    inputs.version = 1

    historical_save = post(client, 'save', command(
        patch(row(quantity=1000) | {'hours': '24'}), version=1))
    assert historical_save.status_code == 200, historical_save.text
    original = client.get('/auto_plan/api/plans/revisions/1').json()
    assert original['evaluation']['output_totals']['expected'] == 855

    converted = client.get('/auto_plan/api/inputs').json()
    stale_request = command(restore(), base=1, version=converted['input_version'])
    changed = client.post('/auto_plan/api/inputs/apply', json={
        'base_version': converted['input_version'],
        'changes': {'line_yields': [
            {'month': '2026-10', 'model_code': 'M', 'percent': '80'}]},
    })
    assert changed.status_code == 200, changed.text
    current_version = changed.json()['input_version']

    stale = preview(client, stale_request)
    assert stale.status_code == 409, stale.text
    assert stale.json()['input_version'] == current_version

    assert client.get('/auto_plan/api/plans/revisions/1').json() == original
    book = workbook(client.get('/auto_plan/api/plans/revisions/1/export.xlsx'))
    exported = json.loads(''.join(row['json_chunk'] for row in records(book, '_record_json')))
    assert exported == original

    request = command(restore(), base=1, version=current_version)
    restored = preview(client, request)
    assert restored.status_code == 200, restored.text
    assert restored.json()['evaluation']['output_totals']['expected'] == 720
    restored_yields = {(row['month'], row['model_code']): row['percent']
                       for row in restored.json()['snapshot']['inputs']['parameters']['line_yields']}
    assert {key: restored_yields[key] for key in (
        ('2026-10', 'M'), ('2026-10', 'N'), ('2026-10', 'P'))} == {
            ('2026-10', 'M'): '80', ('2026-10', 'N'): '95', ('2026-10', 'P'): '95'}

    restored_save = post(client, 'save', request)
    assert restored_save.status_code == 200, restored_save.text
    saved_revision = restored_save.json()['revision']
    current = client.get(f'/auto_plan/api/plans/revisions/{saved_revision}').json()
    assert current['evaluation']['output_totals']['expected'] == 720
    assert current['lineage']['history_restore'] == {'kind': 'revision', 'id': 1}
    assert client.get('/auto_plan/api/plans/revisions/1').json() == original

    adopted = post(client, 'adopt', command(
        base=saved_revision, version=current_version, adoption=0))
    assert adopted.status_code == 200, adopted.text
    assert adopted.json()['latest_adoption'] == 1


def test_restore_resolves_server_reference_into_new_draft_with_current_inputs_and_lineage(shared_client):
    client, plans, _ = shared_client
    assert post(client, 'save', command(patch(row()))).status_code == 200
    old = client.get('/auto_plan/api/plans/revisions/1').json()
    assert post(client, 'save', command(patch(row(quantity=125)), base=1)).status_code == 200
    preferences = {'kind': 'set_preferences', 'preferences': {'near_weeks': 2, 'near_week_weight': '7'}}
    assert post(client, 'save', command(preferences, base=2)).status_code == 200
    request = command(restore(), base=3, version=2)
    response = preview(client, request)
    assert response.status_code == 200, response.text
    proposed = response.json()
    assert proposed['plan']['segments'][0]['quantity'] == 100
    assert proposed['snapshot']['preferences']['near_week_weight'] == '7'
    assert proposed['snapshot']['inputs']['input_version'] == 2
    assert proposed['snapshot']['time']['as_of'] == '2026-10-14T00:00:00Z'
    saved = post(client, 'save', request)
    assert saved.status_code == 200, saved.text
    assert (saved.json()['revision'], saved.json()['latest_adoption']) == (4, 0)
    assert post(client, 'save', request).json() == saved.json()
    assert client.get('/auto_plan/api/plans/revisions/1').json() == old
    restored = client.get('/auto_plan/api/plans/revisions/4').json()
    assert restored['operations'] == [restore()]
    assert restored['lineage']['history_restore'] == {'kind': 'revision', 'id': 1}


def test_restore_rejects_current_quantity_locks_without_copying_old_context(shared_client):
    client, _, _ = shared_client
    assert post(client, 'save', command(patch(row()))).status_code == 200
    assert post(client, 'save', command(patch(row(quantity=125)), base=1)).status_code == 200
    lock = {'kind': 'lock_cells', 'cells': [{'date': '2026-10-15', 'line_code': 'L', 'model_code': 'M', 'quantity': 125}]}
    assert post(client, 'save', command(lock, base=2)).status_code == 200
    blocked = preview(client, command(restore(), base=3, version=2))
    assert blocked.status_code == 409, blocked.text
    assert blocked.json()['conflicts'][0]['code'] == 'locked_cell'
    assert client.get('/auto_plan/api/plans/current').json()['plan']['segments'][0]['quantity'] == 125


def test_restore_has_inclusive_30_day_boundary_and_adoption_is_permanent(shared_client):
    client, _, _ = shared_client
    assert post(client, 'save', command(patch(row()))).status_code == 200
    adopted = post(client, 'adopt', command(base=1))
    assert adopted.status_code == 200, adopted.text
    created = datetime(2026, 10, 14, tzinfo=timezone.utc)
    client.app.state.clock = lambda: created + timedelta(days=30)
    assert client.post('/auto_plan/api/auth/login', json={'username': 'EDITOR', 'password': 'company-secret'}).status_code == 200
    at_boundary = preview(client, command(restore(), base=2, adoption=1))
    assert at_boundary.status_code == 200, at_boundary.text
    client.app.state.clock = lambda: created + timedelta(days=30, microseconds=1)
    expired = preview(client, command(restore(), base=2, adoption=1))
    assert expired.status_code == 410, expired.text
    permanent = preview(client, command(restore(1, 'adoption'), base=2, adoption=1))
    assert permanent.status_code == 200, permanent.text
    listed = client.get('/auto_plan/api/plans/history').json()
    assert [(r['revision'], r['adoption']) for r in listed['items']] == [(2, 1)]
    assert client.get('/auto_plan/api/plans/revisions/1').status_code == 200


def workbook(response):
    from io import BytesIO
    from openpyxl import load_workbook
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content), data_only=False)


def records(book, name):
    values = list(book[name].values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_selected_excel_is_lossless_reproducible_and_deduplicates_only_attribution_group_totals(shared_client):
    import hashlib
    import json
    from auto_plan.input_models import InputSnapshot
    from auto_plan.plan_models import EvaluationSnapshot
    from auto_plan.rules import evaluate
    from tests.test_output_preview import output_snapshot, confirmation
    client, _, inputs = shared_client
    data = output_snapshot(1072, '2026-10-15')
    data['inputs']['business']['demands'] = [{'model_code': 'M', 'week': '2026-10-12', 'quantity': 600}, {'model_code': 'M', 'week': '2026-10-19', 'quantity': 400}]
    data['inputs']['business']['confirmations'] = [confirmation('material', '2026-10-19', 429, '2026-10-15', batch='=MAT')]
    inputs.versions[1] = InputSnapshot.model_validate(data['inputs'] | {'input_version': 1})
    segment = data['plan']['segments'][0] | {'hours': '60001/20000'}
    assert post(client, 'save', command(patch(segment))).status_code == 200
    selected = client.get('/auto_plan/api/plans/revisions/1').json()
    book = workbook(client.get('/auto_plan/api/plans/revisions/1/export.xlsx'))
    assert records(book, 'production')[0]['hours_exact'] == '60001/20000'
    assert records(book, 'production')[0]['hours_display'] == '3.0001'
    assert [r['quantity_exact'] for r in records(book, 'output_portions')] == ['500', '500']
    assert [(r['on_time_exact'], r['gap_exact']) for r in records(book, 'weekly_results')] == [('600', '0'), ('400', '0')]
    assert [r['apportioned_input_exact'] for r in records(book, 'attribution_groups')] == ['643', '429']
    assert len(records(book, 'attributions')) == 3
    assert sum(int(r['apportioned_input_exact']) for r in records(book, 'attribution_groups')) == 1072
    assert records(book, 'confirmation_usage')[0]['used_exact'] == '429'
    assert records(book, 'material_batches')[0]['batch_id'] == '=MAT'
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
    raw = ''.join(row['json_chunk'] for row in records(book, '_record_json'))
    assert json.loads(raw) == selected
    assert records(book, '_manifest')[0]['sha256'] == hashlib.sha256(raw.encode()).hexdigest()
    restored = json.loads(raw)
    assert evaluate(EvaluationSnapshot.model_validate(restored['snapshot'])).model_dump(mode='json') == restored['evaluation']
    assert {'capacities', 'model_yields', 'cycle_times', 'confirmations', 'calendar', 'execution', 'reset', 'locks', 'preferences', 'switch_allocations', 'recipe_hours', 'current_day_proof', 'prior_recipe_hours', 'prior_switches'} <= set(book.sheetnames)


def test_restore_protects_newly_recorded_actuals_and_rechecks_latest_input_on_adoption(shared_client):
    from tests.test_execution_reset import current_command, fact
    client, _, _ = shared_client
    assert post(client, 'save', command(patch(row(day='2026-10-14')))).status_code == 200
    assert post(client, 'save', current_command(client, patch(row('actual', day='2026-10-14')),
        {'kind':'record_execution','rows':[fact(150)],'cutoff':'2026-10-14T08:00:00+08:00'})).status_code == 200
    rejected = preview(client, current_command(client, restore()))
    assert rejected.status_code == 409, rejected.text
    assert any(r['code'] == 'executed_segment' for r in rejected.json()['conflicts'])
    request = current_command(client, restore(2))
    assert preview(client, request).status_code == 200
    prefs={'kind':'set_preferences','preferences':{'near_week_weight':'9'}}
    assert post(client, 'save', current_command(client, prefs)).status_code == 200
    stale = post(client, 'adopt', request)
    assert stale.status_code == 409, stale.text
    assert stale.json()['input_version'] == 3
    assert client.get('/auto_plan/api/plans/current').json()['latest_adoption'] == 0


def test_restore_keeps_current_reset_source_and_does_not_copy_historical_lineage(shared_client):
    from tests.test_execution_reset import seed_reset_source, reset_operation, current_command
    client, plans, _ = shared_client
    seed_reset_source(client, plans)
    assert post(client, 'adopt', current_command(client, reset_operation())).status_code == 200
    before = client.get('/auto_plan/api/plans/current').json()
    checked = preview(client, current_command(client, restore(1, 'adoption')))
    assert checked.status_code == 200, checked.text
    assert checked.json()['context']['reset'] == before['context']['reset']
    assert checked.json()['context']['lineage']['resets'] == before['context']['lineage']['resets']
    assert checked.json()['snapshot']['baseline']['version'] == 2


def test_viewer_can_select_and_export_but_cannot_restore_save_or_adopt_and_revocation_stops_reads(shared_client):
    client, _, _ = shared_client
    assert post(client, 'adopt', command(patch(row()))).status_code == 200
    access=client.app.state.access_store
    access.users['READER']={'user_id':'READER','roles':['viewer'],'active':True}
    assert client.post('/auto_plan/api/auth/login',json={'username':'READER','password':'company-secret'}).status_code == 200
    assert client.get('/auto_plan/api/plans/history').status_code == 200
    assert client.get('/auto_plan/api/plans/adoptions/1').status_code == 200
    assert client.get('/auto_plan/api/plans/adoptions/1/export.xlsx').status_code == 200
    assert preview(client, command(restore(1,'adoption'),base=1,adoption=1)).status_code == 403
    assert post(client,'save',command(restore(),base=1,adoption=1)).status_code == 403
    assert post(client,'adopt',command(restore(),base=1,adoption=1)).status_code == 403
    access.users['READER']['active']=False
    assert client.get('/auto_plan/api/plans/history').status_code == 401
    assert client.get('/auto_plan/api/plans/adoptions/1/export.xlsx').status_code == 401


def test_restore_rechecks_expiration_on_save_and_history_paginates(shared_client):
    client, _, _ = shared_client
    for index in range(3):
        assert post(client,'save',command(patch(row(quantity=100+index)),base=index)).status_code == 200
    page=client.get('/auto_plan/api/plans/history?limit=2').json()
    assert [r['revision'] for r in page['items']] == [3,2]
    assert page['next_before'] == 2
    assert [r['revision'] for r in client.get('/auto_plan/api/plans/history?limit=2&before=2').json()['items']] == [1]
    created=datetime(2026,10,14,tzinfo=timezone.utc)
    client.app.state.clock=lambda:created+timedelta(days=30)
    client.post('/auto_plan/api/auth/login',json={'username':'EDITOR','password':'company-secret'})
    request=command(restore(),base=3)
    assert preview(client,request).status_code == 200
    client.app.state.clock=lambda:created+timedelta(days=30,microseconds=1)
    assert post(client,'save',request).status_code == 410
    assert client.get('/auto_plan/api/plans/current').json()['revision'] == 3


def test_excel_preserves_current_day_joint_proof_prior_switch_hours_and_exact_counts(shared_client):
    from auto_plan.input_models import InputSnapshot
    from tests.test_execution_reset import current_command, fact
    client, _, inputs=shared_client
    data=inputs.read().model_dump(mode='json')
    data['parameters']['switch_durations']=[{'from_recipe':'R','to_recipe':'S','hours':'1'}]
    inputs.versions[1]=InputSnapshot.model_validate(data)
    client.app.state.clock=lambda:datetime(2026,10,12,tzinfo=timezone.utc)
    client.post('/auto_plan/api/auth/login',json={'username':'EDITOR','password':'company-secret'})
    planned={'kind':'replace_plan','plan':{'initial_states':[{'line_code':'L','recipe_code':'R'}],
        'segments':[row('s','P','2026-10-12',quantity=1)|{'hours':'3'}]}}
    facts={'kind':'record_execution','rows':[fact(1,'P','2026-10-12')],'cutoff':'2026-10-12T08:00:00+08:00'}
    assert post(client,'save',current_command(client,planned,facts)).status_code == 200
    confirmation={'kind':'confirm_day_capacity','confirmations':[{'line_code':'L','remaining_hours':'0.9','current_recipe':'S'}],'observed_at':'2026-10-12T08:00:00+08:00'}
    saved=post(client,'save',current_command(client,confirmation))
    assert saved.status_code == 200,saved.text
    book=workbook(client.get('/auto_plan/api/plans/revisions/2/export.xlsx'))
    assert records(book,'current_day_proof')[0]['prior_switch_hours_exact'] == '1'
    assert records(book,'prior_switches')[0]['hours_exact'] == '1'
    selected=client.get('/auto_plan/api/plans/revisions/2').json()
    assert records(book,'prior_recipe_hours')[0]['hours_exact'] == selected['snapshot']['plan']['current_day_witnesses'][0]['prior_recipe_hours'][0]['hours']
    assert records(book,'execution_day_capacities')[0]['remaining_hours_exact'] == '9/10'
    huge=9007199254740993
    assert post(client,'save',current_command(client,patch(row('huge',day='2026-10-15',quantity=huge)))).status_code == 200
    book=workbook(client.get('/auto_plan/api/plans/revisions/3/export.xlsx'))
    assert records(book,'production')[1]['quantity_exact'] == '9007199254740993'


def test_excel_keeps_unrepresentable_xml_notes_lossless_in_canonical_record(shared_client):
    import json
    client, _, _=shared_client
    changed=client.post('/auto_plan/api/inputs/apply',json={'base_version':1,'changes':{'confirmations':[
        {'model_code':'M','week':'2026-10-19','kind':'material','batch_id':'notes','status':'pending','note':'现场\x01备注'}]}})
    assert changed.status_code == 200,changed.text
    assert post(client,'save',command(patch(row()),version=2)).status_code == 200
    book=workbook(client.get('/auto_plan/api/plans/revisions/1/export.xlsx'))
    assert records(book,'confirmations')[0]['note'] == '现场\\u0001备注'
    raw=''.join(r['json_chunk'] for r in records(book,'_record_json'))
    assert json.loads(raw)['snapshot']['inputs']['business']['confirmations'][0]['note'] == '现场\x01备注'


def test_excel_marks_long_reason_display_and_retains_full_audited_reason(shared_client):
    import json
    client, _, _=shared_client
    reason='现场复核' * 10000
    action={'kind':'record_execution','rows':[{'date':'2026-10-14','line_code':'L','model_code':'M','quantity':1}],
            'cutoff':'2026-10-14T08:00:00+08:00','reason':reason}
    assert post(client,'save',command(action)).status_code == 200
    book=workbook(client.get('/auto_plan/api/plans/revisions/1/export.xlsx'))
    assert records(book,'execution_audit')[0]['reason'].endswith('[显示已截短；完整值见 _record_json]')
    raw=''.join(r['json_chunk'] for r in records(book,'_record_json'))
    assert json.loads(raw)['snapshot']['execution']['audit'][0]['reason'] == reason
