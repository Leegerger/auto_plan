"""Explicit standard first-import tables. No legacy-workbook inference."""
from io import BytesIO
import json
from openpyxl import load_workbook
from openpyxl.comments import Comment
from .input_tables import HEADERS, workbook_bytes, excel_rows, paste_rows
from .inputs import error
from .input_models import boolean_value

IMPORT_HEADERS=HEADERS|{
    'segments':('id','date','line_code','order','model_code','quantity'),
    'demands':('model_code','week','week_label','quantity'),
    'coverage':('start_date','end_date','line_codes','model_codes','blanks_are_zero'),
    'reset':('batch_id','production_week','cutoff','model_code','wip','net_balance'),
    'execution_facts':('date','line_code','model_code','quantity'),
    'execution_coverage':('start_date','end_date','line_codes','model_codes','blanks_are_zero'),
    'initial_states':('line_code','recipe_code'),
    'settings':('complete_weekly_demands','execution_cutoff')}
DEMAND_HEADERS={'demand_changes':('model_code','week','week_label','delta')}


def template_bytes(demands=False):
    from openpyxl import Workbook
    book=Workbook() if demands else load_workbook(BytesIO(workbook_bytes()))
    if demands: book.remove(book.active)
    for name,headers in (DEMAND_HEADERS if demands else IMPORT_HEADERS).items():
        if name in book.sheetnames: continue
        sheet=book.create_sheet(name);sheet.append(headers);sheet.freeze_panes='A2'
        for cell in sheet[1]:
            sheet.column_dimensions[cell.column_letter].width=26
            cell.comment=Comment('必须按标准列填写。日期 YYYY-MM-DD，周别 WYYYYWW。编码列表用 | 分隔，含 | 的编码可使用 JSON 字符串数组。顺序必须人工填写。覆盖范围内空白是否为零须填写 true/false。截止时点须含时区。', 'Auto Plan')
    out=BytesIO();book.save(out);return out.getvalue()


def normalize_tables(tables,errors):
    for kind in ('coverage','execution_coverage'):
        for row in tables.get(kind,[]):
            for key in ('line_codes','model_codes'):
                value=row.get(key)
                if isinstance(value,str):
                    try: row[key]=json.loads(value) if value.startswith('[') else value.split('|')
                    except ValueError: errors.append(error(kind,row['_row'],key,'编码列表格式无效'))
            try: row['blanks_are_zero']=boolean_value(row.get('blanks_are_zero'))
            except ValueError as exc: errors.append(error(kind,row['_row'],'blanks_are_zero',str(exc)))
    for row in tables.get('reset',[]):
        from .demands import signed_integer
        try: row['net_balance']=signed_integer(row.get('net_balance'))
        except ValueError as exc: errors.append(error('reset',row['_row'],'net_balance',str(exc)))
    return tables


def bundle(tables,errors,tokens):
    tables=normalize_tables(tables,errors)
    settings=tables.get('settings',[])
    if len(settings)!=1: errors.append(error('settings',2,'settings','必须填写且仅填写一行完整周需求确认和执行事实截止时点'))
    setting=settings[0] if settings else {}
    try: complete=boolean_value(setting.get('complete_weekly_demands'))
    except ValueError as exc: complete=False;errors.append(error('settings',2,'complete_weekly_demands',str(exc)))
    reset_rows=tables.get('reset',[]); reset={}
    if reset_rows:
        reset={key:reset_rows[0].get(key) for key in ('batch_id','production_week','cutoff')}
        reset['models']=[{key:r.get(key) for key in ('model_code','wip','net_balance')} for r in reset_rows]
        for row in reset_rows:
            if any(row.get(key)!=reset[key] for key in ('batch_id','production_week','cutoff')): errors.append(error('reset',row['_row'],'batch_id','同批周一快照的批号、周一日期和截止时点必须一致'))
    payload=tokens|{'changes':{kind:rows for kind,rows in tables.items() if kind in HEADERS},
        'segments':tables.get('segments',[]),'demands':tables.get('demands',[]),
        'coverage':tables.get('coverage',[]),'complete_weekly_demands':complete,'reset':reset,
        'execution':{'facts':[{k:v for k,v in r.items() if k!='_row'} for r in tables.get('execution_facts',[])],
                     'coverage':[{k:v for k,v in r.items() if k!='_row'} for r in tables.get('execution_coverage',[])],
                     'cutoff':setting.get('execution_cutoff')},
        'initial_states':tables.get('initial_states',[])}
    return payload,errors


def excel_bundle(content,tokens):
    tables,errors=excel_rows(content,IMPORT_HEADERS)
    return bundle(tables,errors,tokens)


def pasted_bundle(texts,tokens):
    tables={};errors=[]
    for name,text in texts.items():
        parsed,found=paste_rows(name,text,IMPORT_HEADERS);tables.update(parsed);errors.extend(found)
    return bundle(tables,errors,tokens)
