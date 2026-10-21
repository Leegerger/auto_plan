"""Standard workbook and tab-separated adapters. Business validation stays in inputs."""
import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from zipfile import ZipFile, BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from .inputs import ROW_TYPES, SECTIONS, error

HEADERS = {kind: tuple(model.model_fields) for kind, model in ROW_TYPES.items()}
MAX_UPLOAD = 10 * 1024 * 1024
MAX_ROWS = 100000


def workbook_bytes(snapshot=None):
    book = Workbook()
    book.remove(book.active)
    document = snapshot.model_dump(mode='json') if snapshot else None
    for kind, headers in HEADERS.items():
        sheet = book.create_sheet(kind)
        sheet.append(headers)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{sheet.cell(1, len(headers)).column_letter}1'
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(1, column)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='315B83')
            sheet.column_dimensions[cell.column_letter].width = 24
            cell.comment = Comment(('允许批次必填，其他状态可留空。' if kind == 'confirmations' and header in {'date','quantity'} else '备注可留空。' if kind == 'confirmations' and header == 'note' else '必填。') + '日期 YYYY-MM-DD；月份 YYYY-MM；百分数填95.50表示95.50%；数量与占用小时为非负整数；CT和切机时间最多两位小数；布尔值 true/false。', 'Auto Plan')
            if header in {'active', 'allowed', 'holiday', 'enabled'}:
                validation = DataValidation(type='list', formula1='"true,false"')
                sheet.add_data_validation(validation)
                validation.add(f'{cell.column_letter}2:{cell.column_letter}100001')
        if document:
            for row in document[SECTIONS[kind]][kind]:
                sheet.append([row[header] for header in headers])
                # Fixed codes/names/decimals remain text even if a name starts with '='.
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
    if snapshot:
        meta = book.create_sheet('_version')
        meta.append(['schema_version', 'input_version'])
        meta.append([snapshot.schema_version, snapshot.input_version])
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def paste_rows(kind, text, headers=HEADERS):
    if kind not in headers:
        return {}, [error(kind, 1, 'sheet', '未知资料类别')]
    reader = csv.reader(StringIO(text), delimiter='\t')
    try:
        rows = list(reader)
    except csv.Error:
        return {}, [error(kind, 1, 'table', '粘贴内容不是有效制表符表格')]
    if not rows or tuple(rows[0]) != headers[kind]:
        return {}, [error(kind, 1, 'headers', '首行列名与标准模板不一致')]
    result, errors = [], []
    if len(rows) > MAX_ROWS + 1:
        return {}, [error(kind, 1, 'table', '单批最多100000行')]
    for index, values in enumerate(rows[1:], 2):
        if not any(value.strip() for value in values):
            continue
        if len(values) != len(headers[kind]):
            errors.append(error(kind, index, 'columns', '列数与标准模板不一致'))
            continue
        result.append(dict(zip(headers[kind], values)) | {'_row': index})
    return {kind: result}, errors


def excel_rows(content, headers=HEADERS):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 200 or sum(entry.file_size for entry in archive.infolist()) > 64 * 1024 * 1024:
                return {}, [error('workbook', 1, 'file', '工作簿解压后超过限制')]
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, ValueError, KeyError, OSError):
        return {}, [error('workbook', 1, 'file', '请上传有效的标准 .xlsx 工作簿')]
    changes, errors, count = {}, [], 0
    try:
        for sheet in book:
            kind = sheet.title
            if kind == '_version':
                continue
            if kind not in headers:
                errors.append(error(kind, 1, 'sheet', '未知工作表，请使用标准模板'))
                continue
            iterator = sheet.iter_rows()
            header_cells = next(iterator, ())
            header_values = tuple(cell.value for cell in header_cells)
            # Formatting an empty extra column cannot alter the input schema.
            while header_values and header_values[-1] is None:
                header_values = header_values[:-1]
            if header_values != headers[kind]:
                errors.append(error(kind, 1, 'headers', '列名或顺序与标准模板不一致'))
                continue
            changes[kind] = []
            for index, cells in enumerate(iterator, 2):
                if not any(cell.value is not None for cell in cells):
                    continue
                count += 1
                if count > MAX_ROWS:
                    return {}, [error(kind, index, 'table', '单批最多100000行')]
                values = {}
                for header, cell in zip(headers[kind], cells):
                    value = cell.value
                    if cell.data_type == 'f':
                        errors.append(error(kind, index, header, '不接受公式，请填写明确值'))
                        value = None
                    elif isinstance(value, (date, datetime)):
                        if isinstance(value, datetime) and value.time().isoformat() != '00:00:00':
                            errors.append(error(kind, index, header, '日期单元格不能包含时间'))
                        value = value.strftime('%Y-%m' if header == 'month' else '%Y-%m-%d')
                    elif (type(value) in {int, float} and header in {'percent', 'days'}) or (type(value) is float and header == 'hours'):
                        number = Decimal(str(value))
                        if header == 'percent' and '%' in cell.number_format:
                            number *= 100
                        value = format(number.normalize(), 'f')
                    values[header] = value
                if any(cell.value is not None for cell in cells[len(headers[kind]):]):
                    errors.append(error(kind, index, 'columns', '存在模板以外的列值'))
                changes[kind].append(values | {'_row': index})
        return changes, errors
    finally:
        book.close()
